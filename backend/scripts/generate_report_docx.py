"""
Utility script to generate formatted Word Document (DOCX) from FINAL_REPORT_CONTENT.md
for university PBL submission.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_report_docx():
    print("--> Generating LexExtract Final Project Report (.docx)...")
    
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styling helper colors
    NAVY = RGBColor(0x1E, 0x3A, 0x8A)
    SLATE = RGBColor(0x33, 0x41, 0x55)
    GOLD = RGBColor(0xD9, 0x77, 0x06)

    # Read markdown source
    root_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    md_path = os.path.join(root_dir, "docs", "FINAL_REPORT_CONTENT.md")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        raw = line.rstrip()
        
        # Title Page H1
        if raw.startswith("# "):
            title = raw.replace("# ", "").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(title)
            run.font.name = "Calibri"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = NAVY
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(12)
            
        # Chapter / Section H2
        elif raw.startswith("## "):
            h2 = raw.replace("## ", "").strip()
            p = doc.add_paragraph()
            run = p.add_run(h2)
            run.font.name = "Calibri"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = NAVY
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            
        # Sub-section H3
        elif raw.startswith("### "):
            h3 = raw.replace("### ", "").strip()
            p = doc.add_paragraph()
            run = p.add_run(h3)
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = SLATE
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            
        # Bullet item
        elif raw.startswith("- "):
            item = raw.replace("- ", "").strip()
            p = doc.add_paragraph(style='List Bullet')
            # Simple bold parsing
            if "**" in item:
                parts = item.split("**")
                for idx, part in enumerate(parts):
                    run = p.add_run(part)
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                    if idx % 2 == 1:
                        run.font.bold = True
            else:
                run = p.add_run(item)
                run.font.name = "Calibri"
                run.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(2)
            
        # Empty line
        elif not raw:
            continue
            
        # Normal text paragraph
        elif not raw.startswith("```") and not raw.startswith("|"):
            p = doc.add_paragraph()
            text = raw
            if "**" in text:
                parts = text.split("**")
                for idx, part in enumerate(parts):
                    run = p.add_run(part)
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                    if idx % 2 == 1:
                        run.font.bold = True
            else:
                run = p.add_run(text)
                run.font.name = "Calibri"
                run.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15

    out_path = os.path.join(root_dir, "docs", "LexExtract_Final_Project_Report.docx")
    doc.save(out_path)
    print(f"--> Saved compiled DOCX to: {out_path} ({os.path.getsize(out_path)} bytes)")

def generate_synopsis_docx():
    print("--> Generating LexExtract Project Synopsis (.docx)...")
    doc = docx.Document()
    NAVY = RGBColor(0x1E, 0x3A, 0x8A)
    SLATE = RGBColor(0x33, 0x41, 0x55)
    
    root_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    md_path = os.path.join(root_dir, "docs", "FINAL_SYNOPSIS.md")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        raw = line.rstrip()
        if raw.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(raw.replace("# ", "").strip())
            run.font.name = "Calibri"
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = NAVY
        elif raw.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(raw.replace("## ", "").strip())
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = NAVY
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif raw.startswith("- ") or raw.startswith("1. ") or raw.startswith("2. "):
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(raw)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
        elif raw:
            p = doc.add_paragraph()
            run = p.add_run(raw)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(4)

    out_path = os.path.join(root_dir, "docs", "LexExtract_Project_Synopsis.docx")
    doc.save(out_path)
    print(f"--> Saved compiled Synopsis DOCX to: {out_path} ({os.path.getsize(out_path)} bytes)")

if __name__ == "__main__":
    generate_report_docx()
    generate_synopsis_docx()
